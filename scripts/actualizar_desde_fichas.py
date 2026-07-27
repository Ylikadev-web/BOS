#!/usr/bin/env python3
"""Extrae país, integración, marca y modelo de fichas técnicas PDF y actualiza el anexo."""

from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

import fitz
from openpyxl import load_workbook

PAIS_MAP = {
    "MEXICO": "MEXICO",
    "MÉXICO": "MEXICO",
    "MEJICO": "MEXICO",
    "CHINA": "CHINA",
    "PRC": "CHINA",
    "ESTADOS UNIDOS": "ESTADOS UNIDOS",
    "EUA": "ESTADOS UNIDOS",
    "USA": "ESTADOS UNIDOS",
    "U.S.A.": "ESTADOS UNIDOS",
}


def normalize_key(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).strip().upper()


def normalize_pais(value: str) -> str:
    key = normalize_key(value)
    for k, v in PAIS_MAP.items():
        if k in key:
            return v
    return key


def pdf_text(path: Path) -> str:
    doc = fitz.open(path)
    parts = [page.get_text("text") for page in doc]
    doc.close()
    return "\n".join(parts)


def partida_from_name(name: str) -> int | None:
    m = re.search(r"PARTIDA\s*(\d+)", name, re.I)
    return int(m.group(1)) if m else None


def find_field(text: str, labels: list[str]) -> str | None:
    for label in labels:
        pattern = rf"{label}\s*[:\-]?\s*([^\n\r]+)"
        m = re.search(pattern, text, re.I)
        if m:
            value = m.group(1).strip()
            value = re.split(r"\s{2,}|\t", value)[0].strip()
            return value
    return None


def parse_ficha(text: str) -> dict:
    pais_raw = find_field(
        text,
        [
            r"PA[IÍ]S\s*DE\s*OR[IÍ]GEN",
            r"OR[IÍ]GEN",
            r"PA[IÍ]S\s*FABRICACI[OÓ]N",
            r"FABRICADO\s*EN",
        ],
    )
    gin_raw = find_field(
        text,
        [
            r"GRADO\s*DE\s*INTEGRACI[OÓ]N\s*NACIONAL(?:\s*\(%?\)?)?",
            r"INTEGRACI[OÓ]N\s*NACIONAL(?:\s*\(%?\)?)?",
            r"GIN(?:\s*\(%?\)?)?",
            r"PORCENTAJE\s*DE\s*INTEGRACI[OÓ]N",
        ],
    )
    marca = find_field(text, [r"MARCA"])
    modelo = find_field(text, [r"MODELO", r"CLAVE", r"SKU", r"C[OÓ]DIGO"])

    gin = None
    if gin_raw:
        m = re.search(r"(\d+(?:\.\d+)?)", gin_raw.replace(",", "."))
        if m:
            gin = float(m.group(1))
            if gin > 1 and gin <= 100:
                gin = int(gin)
            elif 0 < gin <= 1:
                gin = int(gin * 100)

    pais = normalize_pais(pais_raw) if pais_raw else None

    if pais is None:
        if re.search(r"HECHO\s*EN\s*M[EÉ]XICO|MADE\s*IN\s*MEXICO", text, re.I):
            pais = "MEXICO"
            gin = gin if gin is not None else 100
        elif re.search(r"HECHO\s*EN\s*CHINA|MADE\s*IN\s*CHINA", text, re.I):
            pais = "CHINA"
            gin = gin if gin is not None else 0

    return {
        "pais": pais,
        "integracion": gin,
        "marca": marca.strip() if marca else None,
        "modelo": modelo.strip() if modelo else None,
    }


def update_excel(excel_path: Path, fichas_dir: Path, sheet_name: str) -> list[tuple]:
    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    updated = []

    for pdf in sorted(fichas_dir.glob("*.pdf")):
        partida = partida_from_name(pdf.name)
        if partida is None:
            continue
        data = parse_ficha(pdf_text(pdf))
        if not any(data.values()):
            continue

        for row in range(2, ws.max_row + 1):
            if ws.cell(row, 1).value != partida:
                continue
            if data["pais"]:
                ws.cell(row, 5, data["pais"])
            if data["integracion"] is not None:
                ws.cell(row, 6, data["integracion"])
            if data["marca"]:
                ws.cell(row, 7, data["marca"].upper())
            if data["modelo"]:
                ws.cell(row, 8, data["modelo"])
            updated.append((partida, data, pdf.name))
            break

    wb.save(excel_path)
    return updated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("fichas_dir", type=Path)
    parser.add_argument("--sheet", default="GPO 1 MAT ELECTRICO Y ELECTRONI")
    args = parser.parse_args()

    rows = update_excel(args.excel, args.fichas_dir, args.sheet)
    print(f"Actualizadas {len(rows)} partidas desde fichas:")
    for partida, data, name in rows:
        print(f"  #{partida} <- {name}: {data}")


if __name__ == "__main__":
    main()
